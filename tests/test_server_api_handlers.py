import json
import os
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from server.db import Database
from server.web import ServerApplication

EVENTS = [{
    "ID": "e1", "NGAYHOC": "01/09/2026", "GIOBATDAU": 7, "PHUTBATDAU": 0,
    "GIOKETTHUC": 9, "PHUTKETTHUC": 30, "TENHOCPHAN": "Machine learning",
    "TENLOPHOCPHAN": "ML-01", "TENPHONGHOC": "A1-101", "GIANGVIEN": "Dr. X",
    "TIETBATDAU": 1, "TIETKETTHUC": 3, "PHANLOAI": "LICHHOC",
}]


class SyncHandlerTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.db = Database(Path(tmp.name) / "t.db")
        self.addCleanup(self.db.close)
        self.user = self.db.get_or_create_user("a")["id"]
        self.sid = self.db.create_session(self.user)
        self.queued, self.errors, self.limited, self.sent = [], [], [], []
        db = self.db
        self.app = SimpleNamespace(
            database=db,
            sync_engine=SimpleNamespace(
                request_sync=lambda sid, run_id=None: self.queued.append((sid, run_id))),
            _api_v1_require_bearer=lambda handler, scope: {"owner_user_id": self.user},
            _api_v1_owned_session=lambda handler, credential, sid: db.get_session(sid),
            _api_v1_error_response=lambda handler, status, code, message, details=None:
                self.errors.append((status, code, details)),
            _api_v1_rate_limited_response=lambda handler, retry_after: self.limited.append(retry_after),
            _json=lambda handler, status, body, content_type=None: self.sent.append((status, body)),
        )

    def call(self):
        ServerApplication._api_v1_sync_session(self.app, SimpleNamespace(), self.sid)

    def login(self):
        self.db.update_session_credentials(self.sid, "uid", "enc", "fp")

    def test_requires_phenikaa_login(self):
        self.call()
        self.assertEqual(self.errors, [(409, "login_required", None)])
        self.assertEqual(self.queued, [])

    def test_disabled_session_is_rejected(self):
        self.login()
        self.db.update_session_status(self.sid, "disabled")
        self.call()
        self.assertEqual(self.errors[0][:2], (409, "session_disabled"))

    def test_queues_a_run_with_a_real_id(self):
        self.login()
        self.call()
        status, body = self.sent[0]
        self.assertEqual(status, 202)
        run_id = body["data"]["runId"]
        self.assertIsInstance(run_id, int)
        self.assertEqual(self.queued, [(self.sid, run_id)])

    def test_second_call_while_running_conflicts(self):
        self.login()
        self.call()
        run_id = self.sent[0][1]["data"]["runId"]
        self.call()
        self.assertEqual(self.errors[-1][:2], (409, "conflict_running"))
        self.assertEqual(self.errors[-1][2], {"runId": run_id})
        self.assertEqual(len(self.queued), 1)

    def test_cooldown_after_a_finished_run(self):
        self.login()
        self.call()
        run_id = self.sent[0][1]["data"]["runId"]
        self.db.finish_sync_run(run_id, ok=True, events=1)
        self.call()
        self.assertEqual(len(self.limited), 1)
        self.assertGreater(self.limited[0], 0)
        self.assertEqual(len(self.queued), 1)
        self.db._write("UPDATE sync_runs SET finished_at = ? WHERE id = ?",
                       ("2000-01-01T00:00:00+00:00", run_id))
        self.call()
        self.assertEqual(len(self.queued), 2)

    def test_cooldown_tolerates_naive_timestamps(self):
        self.login()
        self.call()
        run_id = self.sent[0][1]["data"]["runId"]
        self.db.finish_sync_run(run_id, ok=True, events=1)
        self.db._write("UPDATE sync_runs SET finished_at = ? WHERE id = ?",
                       ("2000-01-01T00:00:00", run_id))
        self.call()
        self.assertEqual(len(self.queued), 2)


class DeleteHandlerTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = tmp.name

    def make(self, google=True):
        rec = SimpleNamespace(removed=[], disconnected=[], deleted=[], sent=[], errors=[])
        lock = threading.RLock()

        def acquire(sid):
            lock.acquire()
            return lock

        app = SimpleNamespace(
            broker=SimpleNamespace(forget_attempt=lambda sid: None, delete_profile=lambda sid: None),
            database=SimpleNamespace(
                update_session_status=lambda sid, status: None,
                get_google_connection=lambda sid: {"calendar_id": "cal"},
                delete_session=lambda sid: rec.deleted.append(sid),
            ),
            config=SimpleNamespace(exports_dir=Path(self.tmp)),
            google=SimpleNamespace(disconnect=lambda sid: rec.disconnected.append(sid)) if google else None,
            _api_v1_require_bearer=lambda handler, scope: {"owner_user_id": 1},
            _api_v1_owned_session=lambda handler, credential, sid: {"id": sid},
            _acquire_profile_lock=acquire,
            _remove_account_calendar=lambda sid, action, warnings: rec.removed.append(action),
            _json=lambda handler, status, body, content_type=None: rec.sent.append((status, body)),
            _api_v1_error_response=lambda handler, status, code, message, details=None:
                rec.errors.append((status, code)),
        )
        return app, rec

    def delete(self, query="", google=True):
        app, rec = self.make(google)
        handler = SimpleNamespace(path="/api/v1/sessions/sid" + query)
        ServerApplication._api_v1_delete_session(app, handler, "sid")
        return rec

    def test_default_keeps_the_google_calendar(self):
        rec = self.delete()
        self.assertEqual(rec.removed, [])
        self.assertEqual(rec.disconnected, ["sid"])
        self.assertEqual(rec.deleted, ["sid"])
        self.assertEqual(rec.sent[0][0], 200)

    def test_hide_and_delete_are_forwarded(self):
        self.assertEqual(self.delete("?calendar=hide").removed, ["hide"])
        self.assertEqual(self.delete("?calendar=delete").removed, ["delete"])

    def test_invalid_choice_deletes_nothing(self):
        rec = self.delete("?calendar=nuke")
        self.assertEqual(rec.errors, [(400, "invalid_request")])
        self.assertEqual(rec.deleted, [])
        self.assertEqual(rec.disconnected, [])

    def test_without_google_the_session_is_still_deleted(self):
        rec = self.delete(google=False)
        self.assertEqual(rec.deleted, ["sid"])
        self.assertEqual(rec.removed, [])


class XlsxTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        (self.root / "sid").mkdir()
        self.app = SimpleNamespace(config=SimpleNamespace(exports_dir=self.root))

    def ensure(self):
        return ServerApplication._api_v1_ensure_xlsx(self.app, "sid")

    def write_json(self, events):
        path = self.root / "sid" / "calendar.json"
        path.write_text(json.dumps(events, ensure_ascii=False), encoding="utf-8")
        return path

    def test_no_source_means_no_xlsx(self):
        self.assertIsNone(self.ensure())

    def test_empty_or_corrupt_source_means_no_xlsx(self):
        self.write_json([])
        self.assertIsNone(self.ensure())
        (self.root / "sid" / "calendar.json").write_text("{", encoding="utf-8")
        self.assertIsNone(self.ensure())

    def test_creates_the_workbook_and_leaves_no_temp_files(self):
        self.write_json(EVENTS)
        path = self.ensure()
        sheet = load_workbook(path)["Calendar"]
        self.assertEqual((sheet["F2"].value, sheet.max_row), ("Machine learning", 2))
        self.assertEqual(list((self.root / "sid").glob("*.tmp.xlsx")), [])

    def test_regenerates_when_the_source_is_newer(self):
        source = self.write_json(EVENTS)
        path = self.ensure()
        newer = EVENTS + [{**EVENTS[0], "ID": "e2", "NGAYHOC": "02/09/2026"}]
        source.write_text(json.dumps(newer, ensure_ascii=False), encoding="utf-8")
        old = source.stat().st_mtime - 10
        os.utime(path, (old, old))
        self.assertEqual(load_workbook(self.ensure())["Calendar"].max_row, 3)


class UnknownRouteTests(unittest.TestCase):
    def run_route(self, method, path):
        errors = []
        app = SimpleNamespace(
            _api_v1_error_response=lambda handler, status, code, message, details=None:
                errors.append((status, code)))
        getattr(ServerApplication, method)(app, SimpleNamespace(path=path, headers={}))
        return errors

    def test_unknown_get_routes_return_the_json_envelope(self):
        for path in ("/api/v1/nope", "/api/v1/nope/deeper", "/api/v1/sessions/x/unknown"):
            self.assertEqual(self.run_route("handle_get", path), [(404, "not_found")], path)

    def test_unknown_post_route_returns_the_json_envelope(self):
        self.assertEqual(self.run_route("handle_post", "/api/v1/nope"), [(404, "not_found")])


if __name__ == "__main__":
    unittest.main()
